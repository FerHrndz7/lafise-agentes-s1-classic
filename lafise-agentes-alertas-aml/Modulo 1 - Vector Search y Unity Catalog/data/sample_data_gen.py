# Databricks notebook source
# MAGIC %md
# MAGIC # Generador de datos sintéticos — Base de Conocimiento PLD/AML
# MAGIC
# MAGIC Genera un corpus **sintético** de Prevención de Lavado de Dinero (PLD/AML) para el laboratorio de
# MAGIC Vector Search + Unity Catalog, y lo escribe en un **Volumen de Unity Catalog** + tablas Delta.
# MAGIC Lo llama `00_setup`, pero también puede ejecutarse solo.
# MAGIC
# MAGIC **Qué crea:**
# MAGIC
# MAGIC *Documentos no estructurados (Volumen UC)* — para parsear y chunkear:
# MAGIC - `politica_pld_lafise.pdf` — Política PLD/FT con **artículos numerados** (estructura ideal para chunking legal).
# MAGIC - `manual_tipologias.pdf` — Manual de Tipologías de Lavado de Activos (estructuración, pass-through, testaferros…).
# MAGIC - `instructivo_alertas.pdf` — Instructivo de escalamiento y disposición de alertas.
# MAGIC - `catalogo_senales_alerta.docx` — Catálogo de señales de alerta por segmento.
# MAGIC - `preguntas_eval.json` — preguntas de prueba con el hecho esperado (para validar el retrieval).
# MAGIC
# MAGIC *Datos estructurados (tablas Delta)* — para filtros por metadata y como hilo para las próximas sesiones:
# MAGIC - `clientes`, `transacciones`, `alertas`, `narrativas_casos`.
# MAGIC
# MAGIC > ⚠️ **Todo es SINTÉTICO.** Nombres, cuentas y montos son inventados con `Faker`. Nunca PII real de LAFISE.
# MAGIC > Los hechos regulatorios (umbrales, plazos) son **plausibles para Centroamérica** (UAF, ROS/RTE, GAFI)
# MAGIC > pero **ficticios y simplificados** para fines didácticos — no son asesoría de cumplimiento.
# MAGIC >
# MAGIC > **Ejecución independiente:** si corre este notebook solo, primero ejecute en una celda
# MAGIC > `%pip install reportlab python-docx faker` y luego `dbutils.library.restartPython()`.

# COMMAND ----------

dbutils.widgets.text("catalog", "", "Catalog")
dbutils.widgets.text("schema", "pld_aml", "Schema base")
dbutils.widgets.text("volume", "documentos", "Volume")
dbutils.widgets.text("num_clientes", "500", "Num clientes")

# COMMAND ----------

import os, re, json, random

CATALOG = dbutils.widgets.get("catalog").strip()
VOLUME = dbutils.widgets.get("volume")
NUM_CLIENTES = int(dbutils.widgets.get("num_clientes"))

# Mismo criterio de aislamiento que 00_setup: el schema lleva el usuario saneado (pld_aml_<usuario>),
# así funciona tanto con un catálogo por asistente como con un catálogo compartido.
user_email = spark.sql("SELECT current_user()").collect()[0][0]
user_prefix = re.sub(r"[^a-z0-9]", "_", user_email.split("@")[0].lower())
assert CATALOG, "Falta el catálogo. Ejecuta este generador desde 00_setup (o pásale el widget 'catalog')."
SCHEMA = f"{(dbutils.widgets.get('schema').strip() or 'pld_aml')}_{user_prefix}"

spark.sql(f"USE CATALOG `{CATALOG}`")
spark.sql(f"CREATE SCHEMA IF NOT EXISTS `{CATALOG}`.{SCHEMA}")
spark.sql(f"CREATE VOLUME IF NOT EXISTS `{CATALOG}`.{SCHEMA}.{VOLUME}")

VOLUME_PATH = f"/Volumes/{CATALOG}/{SCHEMA}/{VOLUME}"
random.seed(42)
print(f"Destino tablas : {CATALOG}.{SCHEMA}")
print(f"Volumen        : {VOLUME_PATH}")

# COMMAND ----------

# MAGIC %md
# MAGIC ## 1. El corpus documental PLD/AML
# MAGIC
# MAGIC Cada documento se define como (título + secciones). Las secciones de la Política usan **artículos
# MAGIC numerados** — así, cuando chunkeemos preservando la estructura, cada chunk podrá citar su artículo.

# COMMAND ----------

# Política PLD: artículos numerados con hechos verificables (umbrales, plazos).
POLITICA = (
    "politica_pld_lafise", "pdf",
    "Política de Prevención de Lavado de Dinero y Financiamiento al Terrorismo (PLD/FT) — Grupo LAFISE",
    [
        ("Artículo 1. Objeto y marco",
         ["La presente política establece el marco de prevención de lavado de dinero y financiamiento al "
          "terrorismo (PLD/FT) de la institución, alineado con las 40 Recomendaciones del GAFI y la normativa "
          "de la Unidad de Análisis Financiero (UAF). Su cumplimiento es obligatorio para todo colaborador."]),
        ("Artículo 2. Enfoque basado en riesgo",
         ["La institución adopta un enfoque basado en riesgo (EBR): asigna a cada cliente un nivel de riesgo "
          "(bajo, medio, alto) según su perfil, actividad económica, geografía y canal, y calibra la intensidad "
          "de la debida diligencia en función de ese nivel."]),
        ("Artículo 3. Debida diligencia del cliente (DDC)",
         ["Se aplica debida diligencia estándar a clientes de riesgo bajo o medio, y debida diligencia "
          "reforzada (EDD) a clientes de alto riesgo. La identificación del beneficiario final es obligatoria "
          "en todos los casos."]),
        ("Artículo 4. Personas Expuestas Políticamente (PEP)",
         ["Toda relación con una Persona Expuesta Políticamente (PEP) requiere debida diligencia reforzada y "
          "la aprobación expresa de la alta gerencia antes de abrir o mantener la cuenta. El monitoreo de las "
          "PEP es continuo."]),
        ("Artículo 5. Reporte de Transacciones en Efectivo (RTE)",
         ["Toda transacción en efectivo igual o superior a USD 10,000 —individual o acumulada por un mismo "
          "cliente durante un mes calendario— genera un Reporte de Transacción en Efectivo (RTE) que se remite "
          "a la UAF. El fraccionamiento deliberado para evadir este umbral constituye una señal de alerta."]),
        ("Artículo 6. Reporte de Operación Sospechosa (ROS)",
         ["Cuando el análisis confirma indicios de una operación sospechosa, el Oficial de Cumplimiento remite "
          "un Reporte de Operación Sospechosa (ROS) a la UAF dentro de un plazo máximo de 3 días hábiles "
          "contados desde la confirmación. El ROS es confidencial: está prohibido informar al cliente (tipping-off)."]),
        ("Artículo 7. Conservación de registros",
         ["Los registros de identificación de clientes, transacciones y análisis de alertas se conservan por un "
          "plazo mínimo de 5 años contados desde el cierre de la relación o la fecha de la operación."]),
        ("Artículo 8. Monitoreo transaccional",
         ["El monitoreo transaccional es automatizado y basado en reglas y tipologías. Genera alertas que el "
          "área de cumplimiento analiza y dispone según el Instructivo de Escalamiento y Disposición de Alertas."]),
        ("Artículo 9. Listas de sanciones",
         ["La institución coteja clientes y contrapartes contra listas de sanciones (OFAC, Consejo de Seguridad "
          "de la ONU y listas locales). Una coincidencia (hit) confirmada bloquea la operación y se escala de "
          "inmediato; nunca se cierra sin acción."]),
    ],
)

TIPOLOGIAS = (
    "manual_tipologias", "pdf",
    "Manual de Tipologías de Lavado de Activos — Cumplimiento LAFISE",
    [
        ("Tipología 1. Estructuración (pitufeo)",
         ["Descripción: fraccionar una suma grande en múltiples depósitos o transferencias menores al umbral de "
          "USD 10,000 para evadir el Reporte de Transacción en Efectivo.",
          "Señales de alerta: múltiples depósitos en efectivo por montos apenas por debajo de USD 10,000; uso de "
          "varias sucursales el mismo día; depósitos por terceros en la cuenta de una misma persona."]),
        ("Tipología 2. Movimiento rápido de fondos (pass-through)",
         ["Descripción: fondos que ingresan y salen de una cuenta en un lapso muy corto (menos de 48 a 72 horas), "
          "dejando un saldo cercano a cero, sin una razón económica aparente.",
          "Señales de alerta: la cuenta funciona como puente; entradas y salidas del mismo monto; beneficiario "
          "final distinto del titular."]),
        ("Tipología 3. Uso de terceros y testaferros",
         ["Descripción: uso de terceros (familiares, empleados, sociedades pantalla) para ocultar al verdadero "
          "beneficiario final de los fondos.",
          "Señales de alerta: cuentas de personas sin capacidad económica que reciben montos altos; poderes "
          "amplios a terceros; sociedades recién constituidas sin actividad real."]),
        ("Tipología 4. Operaciones inconsistentes con el perfil",
         ["Descripción: transacciones cuyo monto, frecuencia o naturaleza no corresponden con el perfil "
          "transaccional declarado ni con la actividad económica del cliente.",
          "Señales de alerta: un cliente de perfil bajo que de pronto mueve montos altos; giros al exterior sin "
          "relación con su negocio; cambios abruptos de comportamiento."]),
        ("Tipología 5. Geografías de alto riesgo",
         ["Descripción: operaciones con origen o destino en jurisdicciones de alto riesgo o no cooperantes según "
          "el GAFI, o en zonas de frontera asociadas a economías ilícitas.",
          "Señales de alerta: transferencias frecuentes hacia países de alto riesgo; uso de centros financieros "
          "extraterritoriales sin justificación; ruteo innecesario de fondos por múltiples jurisdicciones."]),
    ],
)

INSTRUCTIVO = (
    "instructivo_alertas", "pdf",
    "Instructivo de Escalamiento y Disposición de Alertas — Unidad PLD",
    [
        ("1. Alcance",
         ["Este instructivo define cómo el analista de la Unidad de Cumplimiento analiza y dispone las alertas "
          "generadas por el monitoreo transaccional."]),
        ("2. Plazo de análisis (SLA)",
         ["Cada alerta debe analizarse y disponerse en un plazo máximo de 15 días hábiles desde su generación. "
          "Las alertas de prioridad alta (score de riesgo elevado o hit en lista) se atienden en 24 horas."]),
        ("3. Disposiciones posibles",
         ["(a) Cerrar sin acción: cuando el análisis descarta el indicio y la operación es consistente con el "
          "perfil. Debe documentarse el fundamento.",
          "(b) Escalar a investigación: cuando persisten indicios que requieren profundización por la unidad de "
          "investigaciones.",
          "(c) Reportar: cuando se confirma la sospecha, se genera un ROS a la UAF dentro de 3 días hábiles."]),
        ("4. Reglas duras",
         ["Un hit confirmado en lista de sanciones (OFAC/ONU) nunca se cierra sin acción: se bloquea y se escala "
          "de inmediato. Toda disposición debe citar la política o la tipología que la fundamenta."]),
        ("5. Priorización",
         ["Las alertas se priorizan por score de riesgo, monto involucrado, y presencia de PEP o de geografías de "
          "alto riesgo."]),
    ],
)

DOCS = [POLITICA, TIPOLOGIAS, INSTRUCTIVO]

# Documento DOCX: catálogo de señales de alerta.
SENALES_DOCX = (
    "catalogo_senales_alerta", "docx",
    "Catálogo de Señales de Alerta por Segmento — LAFISE",
    [
        ("1. Banca de personas",
         ["Depósitos en efectivo fraccionados; ingresos incompatibles con la ocupación declarada; múltiples "
          "cuentas sin propósito claro; uso intensivo de terceros."]),
        ("2. Banca de empresas",
         ["Facturación inconsistente con el flujo bancario; sociedades sin sustancia económica; pagos a "
          "proveedores en jurisdicciones de alto riesgo; cambios frecuentes de beneficiario final."]),
        ("3. Comercio exterior",
         ["Sobrefacturación o subfacturación de mercancías; pagos por bienes que no se transportan; triangulación "
          "de pagos entre países no relacionados con la operación comercial."]),
    ],
)

# Preguntas de evaluación: hecho esperado + documento fuente (para validar el retrieval).
PREGUNTAS_EVAL = [
    {"pregunta": "¿Cuál es el umbral para reportar una transacción en efectivo (RTE) a la UAF?",
     "hecho_esperado": "USD 10,000", "fuente": "politica_pld_lafise.pdf"},
    {"pregunta": "¿En qué plazo debe remitirse un ROS a la UAF tras confirmar la sospecha?",
     "hecho_esperado": "3 días hábiles", "fuente": "politica_pld_lafise.pdf"},
    {"pregunta": "¿Por cuánto tiempo se conservan los registros de clientes y transacciones?",
     "hecho_esperado": "5 años", "fuente": "politica_pld_lafise.pdf"},
    {"pregunta": "¿Qué es la estructuración o pitufeo?",
     "hecho_esperado": "fraccionar montos para evadir el umbral de USD 10,000", "fuente": "manual_tipologias.pdf"},
    {"pregunta": "¿Cuál es el plazo (SLA) máximo para analizar y disponer una alerta?",
     "hecho_esperado": "15 días hábiles", "fuente": "instructivo_alertas.pdf"},
    {"pregunta": "¿Qué se debe hacer ante un hit confirmado en lista de sanciones?",
     "hecho_esperado": "bloquear y escalar de inmediato; nunca cerrar sin acción", "fuente": "instructivo_alertas.pdf"},
]

# COMMAND ----------

# MAGIC %md
# MAGIC ## 2. Renderizado a PDF y DOCX en el Volumen (idempotente)

# COMMAND ----------

from io import BytesIO
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from docx import Document as DocxDocument


def render_pdf(title, sections):
    buf = BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=LETTER, topMargin=0.9 * inch, bottomMargin=0.9 * inch)
    styles = getSampleStyleSheet()
    flow = [Paragraph(title, styles["Title"]), Spacer(1, 0.25 * inch)]
    for heading, paragraphs in sections:
        flow.append(Paragraph(heading, styles["Heading2"]))
        for p in paragraphs:
            flow.append(Paragraph(p, styles["BodyText"]))
        flow.append(Spacer(1, 0.15 * inch))
    doc.build(flow)
    return buf.getvalue()


def render_docx(title, sections):
    d = DocxDocument()
    d.add_heading(title, level=0)
    for heading, paragraphs in sections:
        d.add_heading(heading, level=1)
        for p in paragraphs:
            d.add_paragraph(p)
    buf = BytesIO()
    d.save(buf)
    return buf.getvalue()

# COMMAND ----------

written = []
for name, fmt, title, sections in DOCS + [SENALES_DOCX]:
    filename = f"{name}.{fmt}"
    data = render_pdf(title, sections) if fmt == "pdf" else render_docx(title, sections)
    dest = os.path.join(VOLUME_PATH, filename)
    with open(dest, "wb") as f:          # sobrescribe -> idempotente
        f.write(data)
    written.append((filename, len(data)))
    print(f"  ✅ {filename}  ({len(data):,} bytes)")

with open(os.path.join(VOLUME_PATH, "preguntas_eval.json"), "w", encoding="utf-8") as f:
    json.dump(PREGUNTAS_EVAL, f, ensure_ascii=False, indent=2)
print(f"  ✅ preguntas_eval.json  ({len(PREGUNTAS_EVAL)} preguntas)")

# COMMAND ----------

# MAGIC %md
# MAGIC ## 3. Tablas estructuradas sintéticas (Faker)
# MAGIC
# MAGIC `clientes`, `transacciones`, `alertas` y `narrativas_casos`. Las **narrativas de casos** son texto libre:
# MAGIC las indexaremos junto con los documentos para que el retrieval devuelva también precedentes similares.

# COMMAND ----------

from faker import Faker
from pyspark.sql import Row
from datetime import datetime, timedelta

fake = Faker("es_MX")
Faker.seed(42)

PAISES_RIESGO = ["Alto", "Alto", "Medio", "Bajo", "Bajo", "Bajo"]
SEGMENTOS = ["Banca de personas", "Banca de empresas", "Comercio exterior"]
TIPOLOGIAS_CAT = ["Estructuración", "Pass-through", "Testaferros",
                  "Inconsistente con perfil", "Geografía de alto riesgo"]

# --- clientes ---
clientes = []
for i in range(NUM_CLIENTES):
    riesgo = random.choice(PAISES_RIESGO)
    clientes.append(Row(
        cliente_id=f"C{i:05d}",
        nombre=fake.name(),
        segmento=random.choice(SEGMENTOS),
        pais=random.choice(["Nicaragua", "Costa Rica", "Honduras", "Panamá", "Guatemala"]),
        es_pep=random.random() < 0.05,
        nivel_riesgo=riesgo,
        fecha_alta=fake.date_between(start_date="-5y", end_date="today").isoformat(),
    ))
df_clientes = spark.createDataFrame(clientes)
df_clientes.write.mode("overwrite").option("overwriteSchema", "true").saveAsTable(f"`{CATALOG}`.{SCHEMA}.clientes")

# --- transacciones ---
txns = []
for i in range(NUM_CLIENTES * 8):
    cid = f"C{random.randint(0, NUM_CLIENTES-1):05d}"
    monto = round(random.choice([random.uniform(50, 3000), random.uniform(9000, 9999),
                                 random.uniform(100, 50000)]), 2)
    txns.append(Row(
        txn_id=f"T{i:07d}",
        cliente_id=cid,
        fecha=(datetime.today() - timedelta(days=random.randint(0, 180))).date().isoformat(),
        monto_usd=monto,
        canal=random.choice(["Efectivo", "Transferencia", "ACH", "Cheque"]),
        pais_contraparte=random.choice(["Nicaragua", "Panamá", "EE.UU.", "Islas Caimán", "Colombia"]),
    ))
spark.createDataFrame(txns).write.mode("overwrite").option("overwriteSchema", "true") \
    .saveAsTable(f"`{CATALOG}`.{SCHEMA}.transacciones")

# --- alertas ---
alertas = []
for i in range(120):
    alertas.append(Row(
        alerta_id=f"A{i:05d}",
        cliente_id=f"C{random.randint(0, NUM_CLIENTES-1):05d}",
        fecha=(datetime.today() - timedelta(days=random.randint(0, 60))).date().isoformat(),
        tipologia=random.choice(TIPOLOGIAS_CAT),
        score_riesgo=random.randint(30, 99),
        monto_involucrado_usd=round(random.uniform(8000, 120000), 2),
        estado=random.choice(["Abierta", "En análisis", "Cerrada", "Escalada", "Reportada"]),
    ))
spark.createDataFrame(alertas).write.mode("overwrite").option("overwriteSchema", "true") \
    .saveAsTable(f"`{CATALOG}`.{SCHEMA}.alertas")

# --- narrativas_casos (texto libre, precedentes históricos) ---
NARRATIVAS = [
    ("K-2024-018", "Estructuración", "Reportado",
     "El cliente realizó nueve depósitos en efectivo de entre USD 9,200 y USD 9,800 en tres sucursales "
     "distintas durante cinco días. El patrón es consistente con estructuración para evadir el umbral de RTE. "
     "Se confirmó la sospecha y se remitió ROS a la UAF. Fundamento: Tipología 1 y Artículo 5 de la política."),
    ("K-2024-042", "Pass-through", "Escalado",
     "Una cuenta empresarial recibió tres transferencias del exterior por USD 45,000 que salieron hacia una "
     "tercera cuenta en menos de 48 horas, dejando saldo cero. No hay soporte comercial. Movimiento rápido de "
     "fondos; se escaló a investigación para identificar al beneficiario final."),
    ("K-2023-233", "Inconsistente con perfil", "Cerrado",
     "Un cliente de banca de personas con perfil bajo recibió USD 30,000. Tras el análisis se documentó que "
     "correspondía a la venta de un inmueble con escritura verificable. Operación consistente; se cerró sin acción."),
    ("K-2024-107", "Geografía de alto riesgo", "Reportado",
     "Transferencias frecuentes hacia una jurisdicción de alto riesgo listada por el GAFI, sin relación con la "
     "actividad declarada del cliente. Se confirmó la sospecha y se reportó a la UAF."),
    ("K-2024-151", "Testaferros", "Escalado",
     "Depósitos altos en la cuenta de una persona sin capacidad económica aparente, con poder amplio otorgado a "
     "un tercero. Posible uso de testaferro; se escaló a investigación."),
]
casos = [Row(caso_id=c[0], tipologia=c[1], disposicion=c[2], narrativa=c[3]) for c in NARRATIVAS]
spark.createDataFrame(casos).write.mode("overwrite").option("overwriteSchema", "true") \
    .saveAsTable(f"`{CATALOG}`.{SCHEMA}.narrativas_casos")

print("✅ Tablas creadas: clientes, transacciones, alertas, narrativas_casos")

# COMMAND ----------

# MAGIC %md
# MAGIC ## 4. Verificación

# COMMAND ----------

print("=" * 64)
print(f"Corpus + tablas en: {CATALOG}.{SCHEMA}")
print("=" * 64)
print("Documentos en el Volumen:")
for f in dbutils.fs.ls(VOLUME_PATH):
    print(f"  {f.name:38s} {f.size:>10,} bytes")
print()
for t in ["clientes", "transacciones", "alertas", "narrativas_casos"]:
    n = spark.table(f"`{CATALOG}`.{SCHEMA}.{t}").count()
    print(f"  tabla {t:20s} {n:>8,} filas")
